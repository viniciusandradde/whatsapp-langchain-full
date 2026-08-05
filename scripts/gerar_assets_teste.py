#!/usr/bin/env python3
"""Gera os arquivos de exemplo que a bateria E2E manda pelo webhook.

Os arquivos ficam versionados em `tests/assets/` — junto de `sample.png`,
`sample.ogg` e `sample.pdf`, que já eram commitados. Este script existe para que
eles sejam **reproduzíveis**: quem precisar mexer no conteúdo roda de novo, em
vez de tentar editar um binário.

Cada arquivo carrega uma SENTINELA no texto (`VSA-E2E-XLSX` e afins). O teste
assere a sentinela, não "veio algum texto" — assim uma extração que devolva
lixo, ou que devolva o arquivo errado, reprova.

Uso:
    uv run python scripts/gerar_assets_teste.py

O `.doc` (Word 97-2003) NÃO é gerado aqui: é formato binário OLE e não há
gerador puro-Python confiável. Ele nasce de uma conversão única do `.docx`,
descrita no fim deste arquivo, e depois vive no repositório como os demais.
"""

from __future__ import annotations

import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
ASSETS = RAIZ / "tests" / "assets"

# O teto do extrator é 10 MB (`file_extractor.MAX_FILE_SIZE_BYTES`). O PDF
# grande precisa passar disso com folga para o teste do teto não ficar na
# fronteira e virar intermitente.
ALVO_PDF_GRANDE_MB = 13


def gerar_docx() -> Path:
    """DOCX de exemplo — e a origem do `.doc`, por conversão.

    O corpo é propositalmente longo: convertido para Word 97-2003, um documento
    curto faz o `antiword` recusar com "the text stream of this file is too
    small to handle". Descoberto na prática, com um `.docx` de três linhas.
    """
    from docx import Document

    caminho = ASSETS / "sample.docx"
    doc = Document()
    doc.add_heading("Proposta comercial de teste", level=1)
    doc.add_paragraph("Sentinela: VSA-E2E-DOCX")
    doc.add_paragraph(
        "Documento gerado por scripts/gerar_assets_teste.py para a bateria E2E "
        "de leitura de documentos. O valor total da proposta é R$ 1.234,56."
    )
    for i in range(1, 41):
        doc.add_paragraph(
            f"Item {i:02d} do escopo — instalação, configuração e homologação "
            "dos equipamentos de infraestrutura, com acompanhamento técnico "
            "durante todo o período contratado e relatório ao final de cada "
            "etapa. Prazo estimado de execução: dez dias úteis."
        )
    doc.save(caminho)
    return caminho


def gerar_xlsx() -> Path:
    """Planilha — o formato que virava `doc.bin` e era recusado.

    Duas abas de propósito: o extrator monta uma seção por aba, e o teste
    confere que a segunda não se perdeu.
    """
    from openpyxl import Workbook

    caminho = ASSETS / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Orcamento"
    ws.append(["Item", "Quantidade", "Valor"])
    ws.append(["Sentinela VSA-E2E-XLSX", 1, 1234.56])
    ws.append(["Consultoria", 10, 500.00])

    ws2 = wb.create_sheet("Resumo")
    ws2.append(["Total", 6234.56])

    wb.save(caminho)
    return caminho


def pdf_acima_do_teto(base: bytes | None = None) -> bytes:
    """PDF acima do teto de 10 MB, em memória.

    **Não é arquivo versionado de propósito.** Um blob de 13 MB no repositório
    seria o maior arquivo do projeto, para provar uma comparação de inteiros —
    `extract_text` mede `len(raw_bytes)` ANTES de abrir o documento, então o
    peso é a única coisa que importa aqui.

    O enchimento entra como comentário PDF depois do `%%EOF`: o arquivo
    continua sendo o `sample.pdf` legítimo até o trailer, e ninguém precisa
    esperar dez minutos de renderização a cada execução.
    """
    origem = base if base is not None else (ASSETS / "sample.pdf").read_bytes()
    alvo = ALVO_PDF_GRANDE_MB * 1024 * 1024
    if len(origem) >= alvo:
        return origem
    # Comentário PDF (`%`) por linha, para o arquivo seguir sendo texto válido
    # do ponto de vista do formato.
    recheio = b"\n%VSA-E2E-PDF-GRANDE " + b"x" * 4000
    repeticoes = (alvo - len(origem)) // len(recheio) + 1
    return origem + recheio * repeticoes


DOC_INSTRUCOES = """
O `.doc` (Word 97-2003) precisa de uma conversão única, porque é binário OLE:

    docker run --rm -v "$PWD/tests/assets:/data" \\
      linuxserver/libreoffice:latest \\
      soffice --headless --convert-to doc --outdir /data /data/sample.docx

Depois confira que o resultado é OLE de verdade (deve começar com D0CF11E0):

    head -c 4 tests/assets/sample.doc | xxd -p

E que o extrator o lê (precisa de `antiword` — está nas imagens Docker):

    uv run python -c "import asyncio,pathlib; \\
      from whatsapp_langchain.shared.file_extractor import extract_text; \\
      print(asyncio.run(extract_text('sample.doc', \\
        pathlib.Path('tests/assets/sample.doc').read_bytes()))[:120])"
"""


def main() -> int:
    if not ASSETS.is_dir():
        print(f"ERRO: {ASSETS} não existe", file=sys.stderr)
        return 1

    for gerar in (gerar_docx, gerar_xlsx):
        caminho = gerar()
        tamanho = caminho.stat().st_size
        print(f"  {caminho.relative_to(RAIZ)}  ({tamanho / 1024:.0f} KB)")

    grande = pdf_acima_do_teto()
    print(
        f"  (em memória) PDF acima do teto: {len(grande) / 1024 / 1024:.1f} MB "
        "— montado pelo teste, não versionado"
    )

    doc = ASSETS / "sample.doc"
    if doc.exists():
        print(
            f"  {doc.relative_to(RAIZ)}  ({doc.stat().st_size / 1024:.0f} KB) — já existe"
        )
    else:
        print("\n.doc ausente:")
        print(DOC_INSTRUCOES)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
